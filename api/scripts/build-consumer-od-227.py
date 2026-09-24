"""Merge Master OD + Consumer OD into a 227-question Consumer workbook."""
from openpyxl import load_workbook, Workbook

MASTER = r"C:\Users\VaishnaviSapkal\Downloads\Master OD Template.xlsx"
CONSUMER = r"C:\Users\VaishnaviSapkal\Downloads\Comsumer OD Template.xlsx"
OUT = r"C:\Users\VaishnaviSapkal\Downloads\Consumer OD Master 227.xlsx"

HEADERS = [
    "Topmost category",
    "Middle Category",
    "Parent Category",
    "Category",
    "Question",
    "Tag",
    "Consumer D2C",
    "Consumer Retail",
    "Format (for additional questions)",
    "Attachments Applicable",
]


def yes(value: str) -> bool:
    return str(value or "").strip().upper() in {"Y", "YES", "TRUE", "1"}


def personalize(text: str) -> str:
    value = str(text or "")
    value = value.replace("<<Company's>>", "KNAV's").replace("<<Compamy's>>", "KNAV's")
    value = value.replace("<<Company>>", "KNAV").replace("<<Compamy>>", "KNAV")
    return value.strip()


def load_rows(path: str, require_flag: bool) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h or "").strip() for h in rows[0]]
    out = []
    for raw in rows[1:]:
        row = {
            headers[i]: ("" if raw[i] is None else str(raw[i]).strip())
            for i in range(len(headers))
        }
        question = personalize(row.get("Question", ""))
        if not question:
            continue
        if require_flag and not (
            yes(row.get("Consumer D2C", "")) or yes(row.get("Consumer Retail", ""))
        ):
            continue
        row["Question"] = question
        row["Attachments Applicable"] = "Y"
        if not row.get("Consumer D2C"):
            row["Consumer D2C"] = "Y"
        if not row.get("Consumer Retail"):
            row["Consumer Retail"] = "Y"
        out.append(row)
    return out


def main() -> None:
    master = load_rows(MASTER, require_flag=True)
    consumer = load_rows(CONSUMER, require_flag=False)
    seen = set()
    merged = []
    for row in master + consumer:
        key = row["Question"].lower()
        if key in seen:
            continue
        seen.add(key)
        merged.append(row)

    wb = Workbook()
    ws = wb.active
    ws.title = "Consumer OD"
    ws.append(HEADERS)
    for row in merged:
        ws.append([row.get(header, "") for header in HEADERS])
    wb.save(OUT)
    print(f"wrote {OUT} with {len(merged)} questions")


if __name__ == "__main__":
    main()
